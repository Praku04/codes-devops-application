import io
import re
import os
import tempfile
import logging
from collections import deque
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image as RLImage,
    PageBreak, Flowable,
)
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.platypus.doctemplate import PageTemplate, BaseDocTemplate
from reportlab.lib.colors import HexColor

# Setup logger
logger = logging.getLogger(__name__)


def _match_mmss_timestamp(text: str, start: int = 0):
    """Match ``\\d{1,2}:\\d{2}(?::\\d{2})?`` at ``text[start:]``. Return (token, end) or None."""
    n = len(text)
    i = start
    if i >= n or not text[i].isdigit():
        return None
    i += 1
    if i < n and text[i].isdigit():
        i += 1
    if i >= n or text[i] != ":":
        return None
    i += 1
    if i + 1 >= n or not text[i].isdigit() or not text[i + 1].isdigit():
        return None
    i += 2
    if i + 2 < n and text[i] == ":" and text[i + 1].isdigit() and text[i + 2].isdigit():
        i += 3
    return text[start:i], i


def _is_figure_heading_line(stripped: str) -> bool:
    """True for ``### Figure N`` (case-insensitive)."""
    if not stripped.lower().startswith("###"):
        return False
    rest = stripped[3:].lstrip()
    if len(rest) < 7 or not rest.lower().startswith("figure"):
        return False
    rest = rest[6:].lstrip()
    if not rest or not rest[0].isdigit():
        return False
    j = 0
    while j < len(rest) and rest[j].isdigit():
        j += 1
    return rest[j:].strip() == ""


def _parse_figure_timestamp_line(stripped: str):
    """
    Parse ``**Timestamp:** MM:SS`` or ``Timestamp: MM:SS`` (optional backticks).
    Returns timestamp string or None.
    """
    low = stripped.lower()
    prefix_star = "**timestamp:**"
    prefix_plain = "timestamp:"
    if low.startswith(prefix_star):
        body = stripped[len(prefix_star) :].strip()
    elif low.startswith(prefix_plain):
        body = stripped[len(prefix_plain) :].strip()
    else:
        return None
    if body.startswith("`"):
        body = body[1:]
        if body.endswith("`"):
            body = body[:-1]
        body = body.strip()
    matched = _match_mmss_timestamp(body, 0)
    if not matched:
        return None
    ts, end = matched
    if body[end:].strip():
        return None
    return ts


def _replace_inline_figure_pipe_timestamps(text: str) -> str:
    """Convert ``**Figure N** | Timestamp: MM:SS`` → ``[IMAGE: MM:SS]`` (linear)."""
    if not text:
        return text
    out: list[str] = []
    i = 0
    lower = text.lower()
    needle = "**figure"
    while i < len(text):
        idx = lower.find(needle, i)
        if idx < 0:
            out.append(text[i:])
            break
        out.append(text[i:idx])
        j = idx + len(needle)
        if j >= len(text) or not text[j].isspace():
            out.append(text[idx])
            i = idx + 1
            continue
        while j < len(text) and text[j].isspace():
            j += 1
        if j >= len(text) or not text[j].isdigit():
            out.append(text[idx])
            i = idx + 1
            continue
        while j < len(text) and text[j].isdigit():
            j += 1
        if not text.startswith("**", j):
            out.append(text[idx])
            i = idx + 1
            continue
        j += 2
        while j < len(text) and text[j].isspace():
            j += 1
        if j >= len(text) or text[j] != "|":
            out.append(text[idx])
            i = idx + 1
            continue
        j += 1
        while j < len(text) and text[j].isspace():
            j += 1
        if not lower.startswith("timestamp:", j):
            out.append(text[idx])
            i = idx + 1
            continue
        j += len("timestamp:")
        while j < len(text) and text[j].isspace():
            j += 1
        matched = _match_mmss_timestamp(text, j)
        if not matched:
            out.append(text[idx])
            i = idx + 1
            continue
        ts, end = matched
        out.append(f"[IMAGE: {ts}]")
        i = end
    return "".join(out)


def _normalize_figure_headers_to_image_tags(md: str) -> str:
    """
    Replace figure/timestamp markdown blocks with ``[IMAGE: MM:SS]`` tags.
    Linear equivalent of the former ReDoS-prone figure regexes (Sonar S5852).
    """
    if not md:
        return md
    lines = md.split("\n")
    out_lines: list[str] = []
    i = 0
    while i < len(lines):
        if _is_figure_heading_line(lines[i].strip()):
            j = i + 1
            while j < len(lines) and lines[j].strip() == "":
                j += 1
            if j < len(lines):
                ts = _parse_figure_timestamp_line(lines[j].strip())
                if ts:
                    out_lines.append(f"[IMAGE: {ts}]")
                    i = j + 1
                    # Former MULTILINE regex \\s*$ could consume a final trailing newline at EOS.
                    if i == len(lines) - 1 and lines[i] == "":
                        i += 1
                    continue
        out_lines.append(lines[i])
        i += 1
    return _replace_inline_figure_pipe_timestamps("\n".join(out_lines))


def _strip_numbered_toc_markdown_links(text: str) -> str:
    """Convert ``1. [Title](#anchor)`` → ``1. Title`` (linear)."""
    if not text or "](" not in text:
        return text
    out: list[str] = []
    i = 0
    n = len(text)
    while i < n:
        if text[i].isdigit():
            j = i
            while j < n and text[j].isdigit():
                j += 1
            if j < n and text[j] == ".":
                k = j + 1
                while k < n and text[k].isspace():
                    k += 1
                if k < n and text[k] == "[":
                    label_start = k + 1
                    label_end = text.find("]", label_start)
                    if label_end > label_start and label_end + 1 < n and text[label_end + 1] == "(":
                        url_end = text.find(")", label_end + 2)
                        if url_end > label_end + 1:
                            out.append(text[i:k])
                            out.append(text[label_start:label_end])
                            i = url_end + 1
                            continue
        out.append(text[i])
        i += 1
    return "".join(out)


def _strip_bullet_toc_markdown_links(text: str) -> str:
    """Convert ``• [Title](#anchor)`` → ``   - Title`` (linear)."""
    if not text or "](" not in text:
        return text
    out: list[str] = []
    i = 0
    n = len(text)
    while i < n:
        ch = text[i]
        if ch in "•·":
            j = i + 1
            while j < n and text[j].isspace():
                j += 1
            if j < n and text[j] == "[":
                label_start = j + 1
                label_end = text.find("]", label_start)
                if label_end > label_start and label_end + 1 < n and text[label_end + 1] == "(":
                    url_end = text.find(")", label_end + 2)
                    if url_end > label_end + 1:
                        out.append("   - ")
                        out.append(text[label_start:label_end])
                        i = url_end + 1
                        continue
        out.append(ch)
        i += 1
    return "".join(out)


def _parse_atx_heading(line: str):
    """
    Parse ``#{1,4} title`` headings. Returns ``(level, title)`` or None.
    Linear replacement for ``^(#{1,4})\\s+(.+?)\\s*$``.
    """
    s = (line or "").strip()
    if not s.startswith("#"):
        return None
    level = 0
    while level < len(s) and level < 4 and s[level] == "#":
        level += 1
    if level == 0:
        return None
    if level < len(s) and s[level] == "#":
        return None
    if level >= len(s) or not s[level].isspace():
        return None
    title = s[level:].strip()
    if not title:
        return None
    return level, title


# -------------------------------------------------------------------------
# Custom DocTemplate for TOC support
# -------------------------------------------------------------------------

class MyDocTemplate(SimpleDocTemplate):
    def afterFlowable(self, flowable):
        "Registers TOC entries."
        if flowable.__class__.__name__ == 'Paragraph':
            text = flowable.getPlainText()
            style = flowable.style.name
            if style == 'Heading1':
                self.notify('TOCEntry', (0, text, self.page))
            elif style == 'Heading2':
                self.notify('TOCEntry', (1, text, self.page))
            elif style == 'Heading3':
                # Optional: Deep linking
                # self.notify('TOCEntry', (2, text, self.page))
                pass


class BookmarkDestination(Flowable):
    """Zero-size anchor: PDF outline (bookmarks pane) + named destination for internal links."""

    def __init__(self, key: str, title: str, level: int = 0):
        Flowable.__init__(self)
        self._key = key
        self._title = (title or "").strip()[:200] or key
        self._level = max(0, min(level, 5))

    def wrap(self, availWidth, availHeight):
        return (0, 0)

    def draw(self):
        self.canv.bookmarkPage(self._key)
        self.canv.addOutlineEntry(self._title, self._key, level=self._level, closed=False)

def extract_text_from_pdf(pdf_file):
    """
    Extracts text from a PDF file object.
    """
    from pypdf import PdfReader
    try:
        reader = PdfReader(pdf_file)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        return f"Error extracting text: {e}"

def extract_text_from_docx(docx_file):
    """
    Extracts text from a DOCX file object.
    """
    from docx import Document
    try:
        doc = Document(docx_file)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        return f"Error extracting text from DOCX: {e}"

def extract_text_from_file(file_object, file_type):
    """
    Extracts text from either PDF or DOCX file.
    
    Args:
        file_object: File-like object (from Streamlit uploader)
        file_type: 'pdf' or 'docx'
    
    Returns:
        Extracted text as string
    """
    if file_type.lower() == 'pdf':
        return extract_text_from_pdf(file_object)
    elif file_type.lower() in ['docx', 'doc']:
        return extract_text_from_docx(file_object)
    else:
        return f"Error: Unsupported file type '{file_type}'"

def create_pdf_from_markdown(output_target, markdown_text, video_path):
    """
    Converts markdown + video frames into a professional SOP-standard PDF.
    Features: cover page, page numbers, section headers/footers, proper figure blocks,
    analysis boxes, and correct image-caption-analysis flow.
    """
    import datetime
    from reportlab.platypus import KeepTogether, HRFlowable, Table, TableStyle, CondPageBreak
    from reportlab.lib import colors

    try:
        # ── Page layout ────────────────────────────────────────────────────
        MARGIN = 0.85 * inch
        PAGE_W, PAGE_H = letter

        # ── Colour palette ─────────────────────────────────────────────────
        C_NAVY    = HexColor("#1a2744")
        C_BLUE    = HexColor("#2563eb")
        C_STEEL   = HexColor("#334155")
        C_GRAY    = HexColor("#64748b")
        C_LGRAY   = HexColor("#f1f5f9")
        C_BORDER  = HexColor("#cbd5e1")
        C_ACCENT  = HexColor("#0ea5e9")
        C_WHITE   = HexColor("#ffffff")
        C_CAPTION = HexColor("#475569")
        C_ANALY   = HexColor("#1e3a5f")
        C_ANALY_BG= HexColor("#eff6ff")
        C_ANALY_BD= HexColor("#bfdbfe")

        # ── Header / Footer callback ────────────────────────────────────────
        doc_title_holder = ["Technical Reference Document"]

        def _header_footer(canvas, doc):
            canvas.saveState()
            w, h = doc.pagesize
            # Header bar
            canvas.setFillColor(C_NAVY)
            canvas.rect(0, h - 0.55*inch, w, 0.55*inch, fill=1, stroke=0)
            canvas.setFillColor(C_WHITE)
            canvas.setFont("Helvetica-Bold", 9)
            title = doc_title_holder[0][:80]
            canvas.drawString(MARGIN, h - 0.35*inch, title)
            canvas.setFont("Helvetica", 8)
            canvas.drawRightString(w - MARGIN, h - 0.35*inch, "CONFIDENTIAL — SOP DOCUMENT")
            # Footer bar
            canvas.setFillColor(C_LGRAY)
            canvas.rect(0, 0, w, 0.45*inch, fill=1, stroke=0)
            canvas.setStrokeColor(C_BORDER)
            canvas.setLineWidth(0.5)
            canvas.line(0, 0.45*inch, w, 0.45*inch)
            canvas.setFillColor(C_GRAY)
            canvas.setFont("Helvetica", 8)
            canvas.drawString(MARGIN, 0.17*inch, f"Generated: {datetime.datetime.now().strftime('%B %d, %Y')}")
            canvas.drawCentredString(w/2, 0.17*inch, "SmartTutor AI — Video Analysis Report")
            canvas.drawRightString(w - MARGIN, 0.17*inch, f"Page {doc.page}")
            canvas.restoreState()

        def _cover_page(canvas, doc):
            canvas.saveState()
            try:
                canvas.setTitle(doc_title_holder[0][:200])
                canvas.setAuthor("SmartTutor AI")
                if hasattr(canvas, "setSubject"):
                    canvas.setSubject("Video analysis — SOP technical reference")
                if hasattr(canvas, "setKeywords"):
                    canvas.setKeywords("SOP, video analysis, technical reference, SmartTutor AI")
            except Exception:
                pass
            w, h = doc.pagesize
            # Full-page navy background
            canvas.setFillColor(C_NAVY)
            canvas.rect(0, 0, w, h, fill=1, stroke=0)
            # Accent stripe
            canvas.setFillColor(C_BLUE)
            canvas.rect(0, h*0.38, w, 4, fill=1, stroke=0)
            # Title
            canvas.setFillColor(C_WHITE)
            canvas.setFont("Helvetica-Bold", 26)
            title = doc_title_holder[0]
            # Word-wrap title
            words = title.split()
            lines, cur = [], []
            for word in words:
                cur.append(word)
                if len(' '.join(cur)) > 38:
                    lines.append(' '.join(cur[:-1]))
                    cur = [word]
            if cur:
                lines.append(' '.join(cur))
            y = h * 0.55
            for ln in lines[:4]:
                canvas.drawCentredString(w/2, y, ln)
                y -= 34
            # Subtitle
            canvas.setFont("Helvetica", 13)
            canvas.setFillColor(C_ACCENT)
            canvas.drawCentredString(w/2, h*0.38 - 28, "Standard Operating Procedure — Technical Reference Document")
            # Date
            canvas.setFont("Helvetica", 11)
            canvas.setFillColor(HexColor("#94a3b8"))
            canvas.drawCentredString(w/2, h*0.38 - 52, datetime.datetime.now().strftime("%B %d, %Y"))
            # Footer
            canvas.setFont("Helvetica", 9)
            canvas.setFillColor(HexColor("#64748b"))
            canvas.drawCentredString(w/2, 0.5*inch, "Generated by SmartTutor AI — Video Analysis System")
            canvas.restoreState()

        doc = BaseDocTemplate(
            output_target,
            pagesize=letter,
            rightMargin=MARGIN, leftMargin=MARGIN,
            topMargin=0.85*inch, bottomMargin=0.65*inch,
        )

        # Two page templates: cover (no header/footer) and body
        from reportlab.platypus.doctemplate import PageTemplate
        from reportlab.platypus.frames import Frame
        cover_frame = Frame(0, 0, PAGE_W, PAGE_H, id='cover')
        body_frame  = Frame(MARGIN, 0.65*inch, PAGE_W - 2*MARGIN, PAGE_H - 1.4*inch, id='body')
        doc.addPageTemplates([
            PageTemplate(id='Cover', frames=[cover_frame], onPage=_cover_page),
            PageTemplate(id='Body',  frames=[body_frame],  onPage=_header_footer),
        ])

        # ── Styles ─────────────────────────────────────────────────────────
        styles = getSampleStyleSheet()

        def _add(name, **kw):
            if name not in styles:
                styles.add(ParagraphStyle(name=name, **kw))

        _add('DocTitle',    parent=styles['Normal'], fontSize=22, leading=28,
             textColor=C_NAVY, fontName='Helvetica-Bold', spaceAfter=6, alignment=TA_CENTER)
        _add('DocSubtitle', parent=styles['Normal'], fontSize=12, leading=16,
             textColor=C_GRAY, fontName='Helvetica', spaceAfter=20, alignment=TA_CENTER)
        _add('H1', parent=styles['Normal'], fontSize=18, leading=24,
             textColor=C_NAVY, fontName='Helvetica-Bold', spaceBefore=18, spaceAfter=10,
             borderPadding=(0,0,4,0))
        _add('H2', parent=styles['Normal'], fontSize=14, leading=20,
             textColor=C_BLUE, fontName='Helvetica-Bold', spaceBefore=14, spaceAfter=8)
        _add('H3', parent=styles['Normal'], fontSize=12, leading=16,
             textColor=C_STEEL, fontName='Helvetica-Bold', spaceBefore=10, spaceAfter=6)
        _add('H4', parent=styles['Normal'], fontSize=11, leading=15,
             textColor=C_GRAY, fontName='Helvetica-Bold', spaceBefore=8, spaceAfter=4)
        _add('Body', parent=styles['Normal'], fontSize=10.5, leading=16,
             textColor=C_STEEL, alignment=TA_JUSTIFY, spaceAfter=8)
        _add('BulletItem', parent=styles['Normal'], fontSize=10.5, leading=15,
             textColor=C_STEEL, leftIndent=18, spaceAfter=5)
        _add('NumItem', parent=styles['Normal'], fontSize=10.5, leading=15,
             textColor=C_STEEL, leftIndent=18, spaceAfter=5)
        _add('TOCEntry1', parent=styles['Normal'], fontSize=11, leading=16,
             textColor=C_NAVY, fontName='Helvetica-Bold', spaceAfter=4)
        _add('TOCEntry2', parent=styles['Normal'], fontSize=10, leading=14,
             textColor=C_STEEL, leftIndent=16, spaceAfter=3)
        _add('FigureLabel', parent=styles['Normal'], fontSize=9, leading=12,
             textColor=C_GRAY, fontName='Helvetica-Bold', alignment=TA_CENTER, spaceAfter=2)
        _add('FigureCaption', parent=styles['Normal'], fontSize=9.5, leading=13,
             textColor=C_CAPTION, fontName='Helvetica-Oblique', alignment=TA_CENTER,
             spaceBefore=4, spaceAfter=8)
        _add('AnalysisLabel', parent=styles['Normal'], fontSize=9, leading=12,
             textColor=C_BLUE, fontName='Helvetica-Bold', spaceAfter=4)
        _add('AnalysisBody', parent=styles['Normal'], fontSize=10, leading=15,
             textColor=C_ANALY, alignment=TA_JUSTIFY, spaceAfter=6)
        _add('BandTitle', parent=styles['Normal'], fontSize=12.5, leading=16,
             textColor=C_WHITE, fontName='Helvetica-Bold', alignment=TA_LEFT, spaceAfter=0)
        _add('CalloutTitle', parent=styles['Normal'], fontSize=10.2, leading=13,
             textColor=C_NAVY, fontName='Helvetica-Bold', spaceAfter=3)
        _add('CalloutBody', parent=styles['Normal'], fontSize=9.5, leading=13,
             textColor=C_STEEL, fontName='Helvetica', alignment=TA_LEFT, spaceAfter=1)
        _add('Italic', parent=styles['Normal'], fontSize=10, leading=14,
             textColor=C_GRAY, fontName='Helvetica-Oblique', spaceAfter=6)
        _add('CodeInline', parent=styles['Normal'], fontSize=9, leading=13,
             fontName='Courier', textColor=HexColor("#be185d"), spaceAfter=6)
        _add('MetaItem', parent=styles['Normal'], fontSize=10.5, leading=15,
             textColor=C_STEEL, leftIndent=0, spaceAfter=5)

        # ── Helpers ────────────────────────────────────────────────────────
        def fmt(text):
            """Markdown inline → ReportLab XML."""
            text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            text = re.sub(r'\*\*\*(.+?)\*\*\*', r'<b><i>\1</i></b>', text)
            text = re.sub(r'\*\*(.+?)\*\*',     r'<b>\1</b>',         text)
            text = re.sub(r'\*([^*]+?)\*',       r'<i>\1</i>',         text)
            text = re.sub(r'`(.+?)`',            r'<font name="Courier" size="9">\1</font>', text)
            return text

        def hr():
            return HRFlowable(width="100%", thickness=0.5, color=C_BORDER,
                              spaceAfter=10, spaceBefore=6)

        def section_band(title_text, level=1):
            """Solid title band for section starts."""
            band_bg = C_NAVY if level == 1 else C_BLUE
            pad_l = 10 if level == 1 else 8
            p = Paragraph(fmt(title_text), styles['BandTitle'])
            tbl = Table([[p]], colWidths=[PAGE_W - 2*MARGIN - 0.06*inch])
            tbl.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), band_bg),
                ('BOX', (0, 0), (-1, -1), 0.8, band_bg),
                ('LEFTPADDING', (0, 0), (-1, -1), pad_l),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 7),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
            ]))
            return [Spacer(1, 6), tbl, Spacer(1, 6)]

        def section_callout(section_title: str):
            """Contextual card for key section types to improve readability."""
            t = (section_title or "").strip().lower()
            if "executive summary" in t:
                title = "Executive Snapshot"
                body = "This section highlights the primary outcomes, business impact, and top-level findings for leadership review."
                bg = HexColor("#ecfeff")
                bd = HexColor("#67e8f9")
            elif "action" in t or "next step" in t or "recommendation" in t:
                title = "Action Checklist"
                body = "Use this section to capture owner-ready actions, timelines, and implementation priorities."
                bg = HexColor("#ecfccb")
                bd = HexColor("#a3e635")
            elif "risk" in t or "mitigation" in t or "security" in t or "compliance" in t:
                title = "Risk & Control Focus"
                body = "This section summarizes key risk signals, compliance considerations, and mitigation controls."
                bg = HexColor("#fff7ed")
                bd = HexColor("#fdba74")
            else:
                return []
            card = Table([[
                Paragraph(title, styles['CalloutTitle']),
            ], [
                Paragraph(body, styles['CalloutBody']),
            ]], colWidths=[PAGE_W - 2*MARGIN - 0.06*inch])
            card.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), bg),
                ('BOX', (0, 0), (-1, -1), 0.9, bd),
                ('LEFTPADDING', (0, 0), (-1, -1), 10),
                ('RIGHTPADDING', (0, 0), (-1, -1), 10),
                ('TOPPADDING', (0, 0), (-1, -1), 7),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
            ]))
            return [card, Spacer(1, 8)]

        def analysis_box(paragraphs):
            """Wrap analysis paragraphs in a styled box."""
            items = []
            for p in paragraphs:
                if p.strip():
                    items.append(Paragraph(fmt(p), styles['AnalysisBody']))
            if not items:
                return []
            label = Paragraph("&#9654; Analysis", styles['AnalysisLabel'])
            tbl = Table([[label]] + [[it] for it in items],
                        colWidths=[PAGE_W - 2*MARGIN - 0.4*inch])
            tbl.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), C_ANALY_BG),
                ('BOX',        (0,0), (-1,-1), 0.8, C_ANALY_BD),
                ('LEFTPADDING',  (0,0), (-1,-1), 12),
                ('RIGHTPADDING', (0,0), (-1,-1), 12),
                ('TOPPADDING',   (0,0), (-1,-1), 8),
                ('BOTTOMPADDING',(0,0), (-1,-1), 8),
                ('ROWBACKGROUNDS',(0,0),(-1,-1),[C_ANALY_BG]),
            ]))
            return [Spacer(1, 6), tbl, Spacer(1, 12)]

        def figure_block(time_str, frame_path, caption_text, analysis_paras, fig_num):
            """Build a KeepTogether figure block: image + caption + analysis."""
            items = [Spacer(1, 8)]
            if frame_path and os.path.exists(frame_path):
                try:
                    img = RLImage(frame_path, width=5.8*inch, height=3.26*inch)
                    items.append(img)
                except Exception:
                    pass
            items.append(Paragraph(f"Figure {fig_num} &nbsp;|&nbsp; Timestamp: {time_str}",
                                   styles['FigureLabel']))
            if caption_text.strip():
                items.append(Paragraph(fmt(caption_text), styles['FigureCaption']))
            items += analysis_box(analysis_paras)
            items.append(hr())
            return KeepTogether(items)

        # ── Normalise markdown ─────────────────────────────────────────────
        md = markdown_text
        # Strip Figure/Timestamp headers (we rebuild them)
        # Support both legacy and current figure markdown formats:
        # 1) **Figure N** | Timestamp: MM:SS
        # 2) ### Figure N + **Timestamp:** `MM:SS` (or Timestamp: MM:SS)
        md = _normalize_figure_headers_to_image_tags(md)
        md = re.sub(r'---+', '', md)
        # Strip TOC markdown links
        md = _strip_numbered_toc_markdown_links(md)
        md = _strip_bullet_toc_markdown_links(md)
        # Remove AI-generated TOC section; we render a clean clickable TOC ourselves.
        def _strip_generated_toc_sections(source_md: str) -> str:
            lines = source_md.split('\n')
            out = []
            skip = False
            toc_level = None
            for ln in lines:
                parsed = _parse_atx_heading(ln)
                if parsed:
                    level, heading_raw = parsed
                    heading_txt = heading_raw.strip().lower()
                    if not skip and heading_txt in {"table of contents", "contents"}:
                        skip = True
                        toc_level = level
                        continue
                    if skip and level <= (toc_level or 2):
                        skip = False
                        toc_level = None
                if not skip:
                    out.append(ln)
            cleaned = '\n'.join(out)
            # Cleanup spacing that can be left after section removal.
            return re.sub(r'\n{3,}', '\n\n', cleaned).strip()

        md = _strip_generated_toc_sections(md)
        # Remove informal phrases
        for ph in [r'[Dd]uring the meeting[,.]?\s*',
                   r'[Ww]as discussed in the meeting[,.]?\s*',
                   r'[Aa]t this point in the meeting[,.]?\s*']:
            md = re.sub(ph, '', md)

        # ── Extract document title ─────────────────────────────────────────
        title_match = re.search(r'^#\s+(.+)', md, re.MULTILINE)
        if title_match:
            doc_title_holder[0] = title_match.group(1).strip()

        # ── Outline / in-document TOC (same order as heading flowables) ────
        def _outline_plain(raw: str) -> str:
            t = re.sub(r'\*\*', '', raw or "")
            t = re.sub(r'[*_`]', '', t)
            return (t.strip()[:200] or "Section")

        def _build_outline_plan(source_md: str):
            plan = []
            n = 0
            first_h1_seen = False
            for raw_line in source_md.split("\n"):
                line = raw_line.strip()
                if line.startswith("#### ") or line.startswith("### "):
                    continue
                if line.startswith("## "):
                    h2 = line[3:].strip()
                    if h2.lower() in {"table of contents", "contents"}:
                        continue
                    n += 1
                    plan.append((2, f"sec{n}", h2))
                elif line.startswith("# "):
                    h1 = line[2:].strip()
                    if h1.lower() in {"table of contents", "contents"}:
                        continue
                    # Avoid listing the document title itself as first TOC entry.
                    if not first_h1_seen:
                        first_h1_seen = True
                        if h1 == doc_title_holder[0]:
                            continue
                    n += 1
                    plan.append((1, f"sec{n}", h1))
            return plan

        outline_plan = _build_outline_plan(md)
        outline_deque = deque(outline_plan)
        outline_auto = [0]

        def _take_outline_entry(expected_level: int):
            if outline_deque and outline_deque[0][0] == expected_level:
                return outline_deque.popleft()
            outline_auto[0] += 1
            return (expected_level, f"_auto{outline_auto[0]}", "")

        # ── Split into segments: text blocks and image markers ─────────────
        IMAGE_RE = r'\[IMAGE:\s*(\d{1,2}:\d{2}(?::\d{2})?)\]'
        raw_parts = re.split(f'({IMAGE_RE})', md)

        # Build ordered list of (type, content) — preserve document order
        segments = []
        i = 0
        while i < len(raw_parts):
            p = raw_parts[i]
            if re.match(IMAGE_RE, p):
                ts = re.match(IMAGE_RE, p).group(1)
                segments.append(('image', ts))
                i += 2  # skip captured group
            else:
                if p.strip():
                    segments.append(('text', p))
                i += 1

        # Same timestamp twice in markdown → same frame twice in PDF; keep first only
        _seen_ts = set()
        _seg_unique = []
        for _seg in segments:
            if _seg[0] == "image":
                if _seg[1] in _seen_ts:
                    continue
                _seen_ts.add(_seg[1])
            _seg_unique.append(_seg)
        segments = _seg_unique

        # ── Temp dir for frames ────────────────────────────────────────────
        temp_frame_dir = tempfile.mkdtemp()
        from utils.video_processor import extract_frame_at_time

        # ── Build story ────────────────────────────────────────────────────
        story = []
        from reportlab.platypus import NextPageTemplate

        # One cover page (default template is Cover), then Body.
        # Do NOT set next template to Cover before the first PageBreak — that
        # was forcing a second full cover page.
        story.append(NextPageTemplate('Body'))
        story.append(PageBreak())

        # Clickable Table of Contents + named destinations (see BookmarkDestination)
        if outline_plan:
            story.append(BookmarkDestination("toc_root", "Table of Contents", 0))
            story.append(Paragraph("Table of Contents", styles['H2']))
            story.append(Spacer(1, 6))
            for lev, key, raw_title in outline_plan:
                indent = "&nbsp;&nbsp;&nbsp;&nbsp;" if lev == 2 else ""
                story.append(Paragraph(
                    f'{indent}<link href="#{key}" color="#2563eb"><u>{fmt(raw_title)}</u></link>',
                    styles['TOCEntry1'] if lev == 1 else styles['TOCEntry2'],
                ))
            story.append(Spacer(1, 14))
            story.append(hr())

        fig_num = 0
        pending_caption = ""
        pending_analysis = []
        in_analysis = False

        def flush_text_block(lines_buf):
            """Convert accumulated text lines to PDF flowables."""
            nonlocal pending_caption, pending_analysis, in_analysis
            for raw_line in lines_buf:
                line = raw_line.strip()
                if not line:
                    if in_analysis and pending_analysis:
                        # End analysis at blank line to avoid bleeding into later content.
                        _flush_analysis()
                    story.append(Spacer(1, 4))
                    continue

                # Headings
                if line.startswith('# '):
                    _flush_analysis()
                    story.append(CondPageBreak(2.4 * inch))
                    ol = _take_outline_entry(1)
                    story.append(BookmarkDestination(ol[1], _outline_plain(ol[2] or line[2:]), 0))
                    story.extend(section_band(line[2:], level=1))
                    story.extend(section_callout(line[2:]))
                    story.append(hr())
                elif line.startswith('## '):
                    _flush_analysis()
                    story.append(CondPageBreak(2.1 * inch))
                    ol = _take_outline_entry(2)
                    story.append(BookmarkDestination(ol[1], _outline_plain(ol[2] or line[3:]), 1))
                    story.extend(section_band(line[3:], level=2))
                    story.extend(section_callout(line[3:]))
                elif line.startswith('### '):
                    _flush_analysis()
                    story.append(CondPageBreak(1.8 * inch))
                    story.append(Paragraph(fmt(line[4:]), styles['H3']))
                elif line.startswith('#### '):
                    _flush_analysis()
                    story.append(CondPageBreak(1.4 * inch))
                    story.append(Paragraph(fmt(line[5:]), styles['H4']))

                # Caption
                elif re.match(r'\*\*Caption:\*\*', line, re.IGNORECASE):
                    _flush_analysis()
                    pending_caption = re.sub(r'\*\*Caption:\*\*\s*', '', line, flags=re.IGNORECASE).strip()

                # Analysis header
                elif re.match(r'\*\*Analysis:\*\*|Analysis:', line, re.IGNORECASE):
                    in_analysis = True
                    rest = re.sub(r'\*\*Analysis:\*\*\s*|Analysis:\s*', '', line, flags=re.IGNORECASE).strip()
                    if rest:
                        pending_analysis.append(rest)

                # Continuation of analysis
                elif in_analysis:
                    pending_analysis.append(line)

                # Bullets
                elif line.startswith('- ') or line.startswith('• '):
                    _flush_analysis()
                    story.append(Paragraph(f"&bull;&nbsp; {fmt(line[2:])}", styles['BulletItem']))
                elif re.match(r'^\d+\.\s', line):
                    _flush_analysis()
                    m = re.match(r'^(\d+)\.\s(.+)', line)
                    if m:
                        story.append(Paragraph(f"{m.group(1)}.&nbsp; {fmt(m.group(2))}", styles['NumItem']))

                # Italic-only line (image placeholder descriptions)
                elif line.startswith('*') and line.endswith('*') and not line.startswith('**'):
                    _flush_analysis()
                    story.append(Paragraph(fmt(line[1:-1]), styles['Italic']))

                # Meta items (key: value)
                elif re.match(r'^\*\*[^*]+:\*\*', line):
                    _flush_analysis()
                    story.append(Paragraph(fmt(line), styles['MetaItem']))

                # Regular paragraph
                else:
                    _flush_analysis()
                    story.append(Paragraph(fmt(line), styles['Body']))

        def _flush_analysis():
            nonlocal pending_analysis, in_analysis
            if in_analysis and pending_analysis:
                story.extend(analysis_box(pending_analysis))
                pending_analysis = []
            in_analysis = False

        # Process segments in document order
        # Caption and Analysis in the markdown come AFTER the [IMAGE: ts] tag,
        # so we need a two-pass approach: when we hit an image, we look ahead
        # in the following text segment for Caption/Analysis before flushing it.
        def _is_structural_line(s: str) -> bool:
            if not s:
                return False
            if s.startswith('#') or s.startswith('---'):
                return True
            if re.match(r'\[IMAGE:\s*\d{1,2}:\d{2}(?::\d{2})?\]', s):
                return True
            if re.match(r'\*\*Figure\s+\d+\*\*', s, re.IGNORECASE):
                return True
            if re.match(r'\*\*Caption:\*\*', s, re.IGNORECASE):
                return True
            if re.match(r'^(Overview|Technical Details|Configuration\s*&\s*Setup|Operational Procedures)\s*:', s, re.IGNORECASE):
                return True
            # Some model outputs use plain subsection titles like "Age Analysis"
            # without markdown heading markers; treat those as structural boundaries.
            if re.match(r'^[A-Z][A-Za-z0-9&/\-]*(\s+[A-Z][A-Za-z0-9&/\-]*){0,6}$', s):
                if s.lower() not in {'overview', 'technical details'}:
                    return True
            return False

        def _looks_like_figure_narrative(s: str) -> bool:
            return bool(re.match(
                r'^(This\s+(figure|image|chart|graph|table|screen)|'
                r'The\s+(figure|image|dashboard|screen|chart|graph|table)|'
                r'(Screenshot|Screen|Dashboard|Chart|Graph|Table)\b)',
                s,
                re.IGNORECASE
            ))

        text_buf = []
        i = 0
        while i < len(segments):
            seg_type, seg_content = segments[i]
            if seg_type == 'text':
                text_buf.extend(seg_content.split('\n'))
                i += 1
            else:
                # Flush accumulated text first (text BEFORE the image)
                flush_text_block(text_buf)
                _flush_analysis()
                text_buf = []

                ts = seg_content

                # Look ahead: the next text segment may contain Caption + Analysis
                # that belong to THIS image
                img_caption = ""
                img_analysis = []
                remaining_lines = []

                if i + 1 < len(segments) and segments[i + 1][0] == 'text':
                    next_lines = segments[i + 1][1].split('\n')
                    in_cap = False
                    in_ana = False
                    for ln in next_lines:
                        stripped = ln.strip()
                        if re.match(r'\*\*Caption:\*\*', stripped, re.IGNORECASE):
                            in_cap = True
                            in_ana = False
                            img_caption = re.sub(r'\*\*Caption:\*\*\s*', '', stripped,
                                                 flags=re.IGNORECASE).strip()
                        elif re.match(r'\*\*Analysis:\*\*|^Analysis:\s*$', stripped, re.IGNORECASE):
                            in_cap = False
                            in_ana = True
                            rest = re.sub(r'\*\*Analysis:\*\*\s*|Analysis:\s*', '', stripped,
                                          flags=re.IGNORECASE).strip()
                            if rest:
                                img_analysis.append(rest)
                        elif in_ana:
                            # Strictly stop this image analysis before next structural marker
                            # to prevent mixing analysis content across sections/images.
                            if (not stripped or _is_structural_line(stripped)):
                                in_ana = False
                                remaining_lines.append(ln)
                            else:
                                img_analysis.append(stripped)
                        elif in_cap and stripped and not re.match(r'\*\*Analysis:\*\*', stripped, re.IGNORECASE):
                            # Multi-line caption continuation
                            img_caption += ' ' + stripped
                            in_cap = False  # caption is one line
                        else:
                            if not in_cap and not in_ana:
                                remaining_lines.append(ln)

                    # Heuristic: if the model did not explicitly mark **Analysis:** but
                    # left a narrative paragraph right after the figure, attach that text
                    # to THIS figure block so users don't see mismatched content later.
                    if remaining_lines:
                        first_idx = None
                        for ridx, rln in enumerate(remaining_lines):
                            rs = rln.strip()
                            if not rs:
                                continue
                            if _is_structural_line(rs):
                                break
                            first_idx = ridx
                            break

                        if first_idx is not None:
                            grab_idx = first_idx
                            grabbed = []
                            while grab_idx < len(remaining_lines):
                                gs = remaining_lines[grab_idx].strip()
                                if not gs:
                                    if grabbed:
                                        break
                                    grab_idx += 1
                                    continue
                                if _is_structural_line(gs):
                                    break
                                grabbed.append(gs)
                                grab_idx += 1

                            if grabbed and (img_caption.strip() or _looks_like_figure_narrative(grabbed[0])):
                                img_analysis.extend(grabbed)
                                del remaining_lines[first_idx:grab_idx]

                    # Replace the next text segment with only the remaining lines
                    if remaining_lines:
                        segments[i + 1] = ('text', '\n'.join(remaining_lines))
                    else:
                        # Remove the consumed text segment
                        segments.pop(i + 1)

                # Build figure block with caption/analysis from the text that followed
                frame_path = extract_frame_at_time(video_path, ts, temp_frame_dir)
                fig_num += 1
                story.append(CondPageBreak(3.9 * inch))
                story.append(figure_block(ts, frame_path, img_caption, img_analysis, fig_num))
                pending_caption = ""
                pending_analysis = []
                in_analysis = False
                i += 1

        # Flush remaining text
        flush_text_block(text_buf)
        _flush_analysis()

        # ── Build PDF ──────────────────────────────────────────────────────
        doc.build(story)
        return True

    except Exception as e:
        logger.error(f"Error creating PDF: {e}", exc_info=True)
        import traceback
        traceback.print_exc()
        return False


def _normalize_video_report_markdown(markdown_text: str) -> tuple[str, str]:
    """Shared markdown cleanup for PDF/DOCX video analysis exports."""
    md = markdown_text
    md = _normalize_figure_headers_to_image_tags(md)
    md = re.sub(r'---+', '', md)
    md = _strip_numbered_toc_markdown_links(md)
    md = _strip_bullet_toc_markdown_links(md)
    title_match = re.search(r'^#\s+(.+)', md, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else "Technical Reference Document"
    return md, title


def _split_video_report_segments(md: str) -> list:
    """Return ordered list of ('text', str) or ('image', timestamp) segments."""
    image_re = r'\[IMAGE:\s*(\d{1,2}:\d{2}(?::\d{2})?)\]'
    raw_parts = re.split(f'({image_re})', md)
    segments = []
    i = 0
    while i < len(raw_parts):
        part = raw_parts[i]
        match = re.match(image_re, part)
        if match:
            segments.append(('image', match.group(1)))
            i += 2
        else:
            if part.strip():
                segments.append(('text', part))
            i += 1
    seen_ts = set()
    unique = []
    for seg in segments:
        if seg[0] == 'image':
            if seg[1] in seen_ts:
                continue
            seen_ts.add(seg[1])
        unique.append(seg)
    return unique


def _strip_inline_markdown(text: str) -> str:
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*([^*]+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text.strip()


def create_docx_from_markdown(output_target, markdown_text: str, video_path: str) -> bool:
    """Convert video-analysis markdown (+ frames) into an editable Word document."""
    import datetime as dt

    try:
        from docx import Document
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from docx.shared import Inches, Pt
        from utils.video_processor import extract_frame_at_time

        md, doc_title = _normalize_video_report_markdown(markdown_text)
        segments = _split_video_report_segments(md)
        temp_frame_dir = tempfile.mkdtemp()

        document = Document()
        title_para = document.add_heading(doc_title, level=0)
        title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        document.add_paragraph(
            f"Generated by SmartTutor AI — {dt.datetime.now().strftime('%B %d, %Y')}"
        ).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        document.add_paragraph()

        fig_num = 0

        def add_text_block(text_blob: str):
            for raw_line in text_blob.split('\n'):
                line = raw_line.strip()
                if not line:
                    continue
                if line.startswith('# '):
                    document.add_heading(_strip_inline_markdown(line[2:]), level=1)
                elif line.startswith('## '):
                    document.add_heading(_strip_inline_markdown(line[3:]), level=2)
                elif line.startswith('### '):
                    document.add_heading(_strip_inline_markdown(line[4:]), level=3)
                elif line.startswith('#### '):
                    document.add_heading(_strip_inline_markdown(line[5:]), level=4)
                elif re.match(r'\*\*Caption:\*\*', line, re.IGNORECASE):
                    continue
                elif re.match(r'\*\*Analysis:\*\*|^Analysis:\s*$', line, re.IGNORECASE):
                    continue
                elif line.startswith('- ') or line.startswith('• '):
                    document.add_paragraph(_strip_inline_markdown(line[2:]), style='List Bullet')
                elif re.match(r'^\d+\.\s', line):
                    m = re.match(r'^\d+\.\s(.+)', line)
                    if m:
                        document.add_paragraph(_strip_inline_markdown(m.group(1)), style='List Number')
                else:
                    document.add_paragraph(_strip_inline_markdown(line))

        idx = 0
        while idx < len(segments):
            seg_type, seg_content = segments[idx]
            if seg_type == 'text':
                add_text_block(seg_content)
                idx += 1
                continue

            ts = seg_content
            img_caption = ""
            img_analysis = []
            if idx + 1 < len(segments) and segments[idx + 1][0] == 'text':
                remaining = []
                in_analysis = False
                for ln in segments[idx + 1][1].split('\n'):
                    stripped = ln.strip()
                    if re.match(r'\*\*Caption:\*\*', stripped, re.IGNORECASE):
                        img_caption = re.sub(
                            r'\*\*Caption:\*\*\s*', '', stripped, flags=re.IGNORECASE
                        ).strip()
                        in_analysis = False
                    elif re.match(r'\*\*Analysis:\*\*|^Analysis:\s*$', stripped, re.IGNORECASE):
                        in_analysis = True
                        rest = re.sub(
                            r'\*\*Analysis:\*\*\s*|Analysis:\s*', '', stripped, flags=re.IGNORECASE
                        ).strip()
                        if rest:
                            img_analysis.append(rest)
                    elif in_analysis:
                        if stripped and not stripped.startswith('#'):
                            img_analysis.append(stripped)
                        else:
                            in_analysis = False
                            remaining.append(ln)
                    else:
                        remaining.append(ln)
                if remaining:
                    segments[idx + 1] = ('text', '\n'.join(remaining))
                else:
                    segments.pop(idx + 1)

            fig_num += 1
            document.add_heading(f'Figure {fig_num} — Timestamp {ts}', level=3)
            frame_path = extract_frame_at_time(video_path, ts, temp_frame_dir)
            if frame_path and os.path.exists(frame_path):
                try:
                    document.add_picture(frame_path, width=Inches(6.0))
                except Exception as img_err:
                    logger.warning("DOCX: could not embed frame at %s: %s", ts, img_err)
            if img_caption:
                cap = document.add_paragraph(img_caption)
                cap.italic = True
            if img_analysis:
                document.add_paragraph('Analysis').runs[0].bold = True
                for para in img_analysis:
                    document.add_paragraph(_strip_inline_markdown(para))
            document.add_paragraph()
            idx += 1

        if hasattr(output_target, 'write'):
            document.save(output_target)
        else:
            document.save(output_target)
        return True
    except Exception as e:
        logger.error("Error creating DOCX: %s", e, exc_info=True)
        return False
